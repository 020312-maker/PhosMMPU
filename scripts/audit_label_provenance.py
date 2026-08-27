"""Trace PhosMMPU functional labels back to the raw Phosphosite workbook.

The processed site table canonicalizes UniProt isoform accessions.  This audit
therefore separates rows genuinely rejected by site validation from rows whose
raw direct ID changed but whose canonical site was retained.
"""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

import pandas as pd


TASKS = (
    ("活性调控", "y_activity", "ON_FUNCTION 包含 activity", "ON_FUNCTION"),
    (
        "分子关联",
        "y_interaction",
        "ON_PROT_INTERACT 有有效内容，或 ON_FUNCTION 包含 molecular association",
        "ON_PROT_INTERACT / ON_FUNCTION",
    ),
    (
        "稳定性/降解",
        "y_proteostasis",
        "ON_FUNCTION 包含 protein stabilization 或 protein degradation",
        "ON_FUNCTION",
    ),
)


def _raw_direct_ids(raw: pd.DataFrame) -> pd.Series:
    accession = raw["ACC_ID"].fillna("").astype(str).str.strip()
    residue = raw["MOD_RSD"].fillna("").astype(str).str.strip()
    return accession + "_" + residue.str.replace("-p", "", regex=False)


def _source_site_map(index: pd.DataFrame, audit: pd.DataFrame) -> pd.DataFrame:
    """Map every source row retained in a canonical site, including duplicates."""
    columns = ["source_row", "site_id"]
    maps = [index.loc[:, columns]]
    if {"source_row", "site_id", "reason"}.issubset(audit.columns):
        duplicate_rows = audit.loc[
            (audit["reason"] == "duplicate_site_id") & audit["site_id"].notna(), columns
        ]
        maps.append(duplicate_rows)
    return pd.concat(maps, ignore_index=True).drop_duplicates(columns)


def _label_rules() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "任务": "活性调控",
                "P 标签规则": "ON_FUNCTION 用不区分大小写的词边界匹配 activity",
                "原始字段": "ON_FUNCTION",
                "0 的含义": "U（未标注），不是确认无功能的负例",
            },
            {
                "任务": "分子关联",
                "P 标签规则": "ON_PROT_INTERACT 去空白后不是 -/空/nan/none，或 ON_FUNCTION 包含 molecular association",
                "原始字段": "ON_PROT_INTERACT；ON_FUNCTION",
                "0 的含义": "U（未标注），不是确认无功能的负例",
            },
            {
                "任务": "稳定性/降解",
                "P 标签规则": "ON_FUNCTION 包含 protein stabilization 或 protein degradation",
                "原始字段": "ON_FUNCTION",
                "0 的含义": "U（未标注），不是确认无功能的负例",
            },
        ]
    )


def build_provenance_tables(
    raw: pd.DataFrame,
    index: pd.DataFrame,
    audit: pd.DataFrame,
    labels: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    """Build all tables needed to explain source IDs, exclusions, and P labels."""
    raw = raw.copy()
    if "source_row" not in raw:
        raw.insert(0, "source_row", range(len(raw)))
    raw["原始直接ID"] = _raw_direct_ids(raw)
    # pandas source_row is zero-based data-row indexing; Excel has one header row.
    raw["原始Excel行号"] = raw["source_row"].astype(int) + 2

    canonical_ids = set(index["site_id"].astype(str))
    raw_ids = set(raw["原始直接ID"].astype(str))
    source_map = _source_site_map(index, audit)
    source_to_site = source_map.set_index("source_row")["site_id"]

    audit_reasons = audit.loc[:, [c for c in ("source_row", "reason") if c in audit.columns]].copy()
    raw_outcomes = raw.merge(audit_reasons, on="source_row", how="left")
    raw_outcomes["canonical_site_id"] = raw_outcomes["source_row"].map(source_to_site)
    raw_outcomes = raw_outcomes.loc[~raw_outcomes["原始直接ID"].isin(canonical_ids)].copy()
    raw_outcomes["处理结论"] = raw_outcomes["reason"].notna().map(
        {True: "excluded", False: "canonicalized_and_retained"}
    )
    raw_outcomes["原因说明"] = raw_outcomes["reason"].fillna(
        "原始 isoform 直接 ID 已规范化为 canonical site_id，位点保留"
    )
    raw_outcomes = raw_outcomes.rename(columns={"reason": "处理原因"})

    canonicalized = index.loc[~index["site_id"].isin(raw_ids)].copy()
    canonicalized = canonicalized.merge(
        raw.loc[:, ["source_row", "原始Excel行号", "ACC_ID", "MOD_RSD", "原始直接ID", "Database", "ON_FUNCTION", "ON_PROT_INTERACT"]],
        on="source_row",
        how="left",
    ).merge(labels, on="site_id", how="left")
    canonicalized = canonicalized.rename(columns={"site_id": "总表site_id"})

    mapped_raw = raw.merge(source_map, on="source_row", how="inner")
    mapped_raw = mapped_raw.merge(labels, on="site_id", how="left")
    positive_rows: list[dict[str, object]] = []
    for _, row in mapped_raw.iterrows():
        for task, column, rule, fields in TASKS:
            if int(row.get(column, 0) or 0) != 1:
                continue
            positive_rows.append(
                {
                    "site_id": row["site_id"],
                    "蛋白": row.get("accession", ""),
                    "原始ACC_ID": row.get("ACC_ID", ""),
                    "MOD_RSD": row.get("MOD_RSD", ""),
                    "原始行号(0基)": row["source_row"],
                    "原始Excel行号": row["原始Excel行号"],
                    "原始直接ID": row["原始直接ID"],
                    "任务": task,
                    "P标签": 1,
                    "生成规则": rule,
                    "触发字段": fields,
                    "Database": row.get("Database", ""),
                    "ON_FUNCTION": row.get("ON_FUNCTION", ""),
                    "ON_PROT_INTERACT": row.get("ON_PROT_INTERACT", ""),
                }
            )
    positive_trace = pd.DataFrame(positive_rows)
    if not positive_trace.empty:
        positive_trace = positive_trace.loc[
            positive_trace["site_id"].isin(set(canonicalized["总表site_id"]))
        ].sort_values(["任务", "site_id", "原始行号(0基)"], ignore_index=True)

    actual_rejections = raw_outcomes.loc[raw_outcomes["处理结论"] == "excluded"]
    rejection_summary = (
        actual_rejections.groupby("处理原因", dropna=False).size().rename("数量").reset_index()
    )
    retained_summary = pd.DataFrame(
        [{"处理原因": "canonicalized_and_retained", "数量": int((raw_outcomes["处理结论"] == "canonicalized_and_retained").sum())}]
    )
    rejection_summary = pd.concat([rejection_summary, retained_summary], ignore_index=True)

    summary = pd.DataFrame(
        [
            ("原始 Phosphosite 行数", len(raw), "原始工作簿的全部数据行"),
            ("总表 canonical 位点数", len(index), "通过校验并按 canonical site_id 合并后的唯一位点"),
            ("直接字符串交集", len(raw_ids & canonical_ids), "原始直接 ID 与总表 site_id 完全相同"),
            ("原始侧直接 ID 未匹配总表", len(raw_outcomes), "包含真正排除与 canonical 化后保留两类"),
            ("其中实际排除", len(actual_rejections), "不进入总表"),
            ("其中 canonical 化后保留", int((raw_outcomes["处理结论"] == "canonicalized_and_retained").sum()), "原始 isoform ID 改为 canonical ID 后进入总表"),
            ("总表侧 canonical 化保留位点", len(canonicalized), "总表有但原始直接字符串没有的位点"),
            ("canonical 化保留的 P 位点", positive_trace["site_id"].nunique() if not positive_trace.empty else 0, "三个任务任一为 P 的唯一位点数"),
            ("canonical 化保留的 P 标签事件", len(positive_trace), "按任务计数；多标签位点会计入多个任务"),
        ],
        columns=["项目", "数量", "说明"],
    ).set_index("项目")

    return {
        "summary": summary,
        "canonicalized": canonicalized,
        "positive_trace": positive_trace,
        "raw_outcomes": raw_outcomes,
        "rejection_summary": rejection_summary,
        "label_rules": _label_rules(),
    }


def write_workbook(tables: dict[str, pd.DataFrame], output: Path) -> None:
    """Write a formatted audit workbook without hiding the raw trace columns."""
    output.parent.mkdir(parents=True, exist_ok=True)
    sheet_names = {
        "summary": "说明",
        "canonicalized": "218条映射",
        "positive_trace": "25个P追溯",
        "raw_outcomes": "4727条去向",
        "rejection_summary": "实际排除分类",
        "label_rules": "P标签规则",
    }
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for key, sheet in sheet_names.items():
            frame = tables[key]
            frame.to_excel(writer, sheet_name=sheet, index=(key == "summary"))

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(output)
    header_fill = PatternFill("solid", fgColor="176B87")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
            cell.fill = header_fill
        for column_cells in worksheet.columns:
            width = max(len(str(cell.value or "")) for cell in column_cells) + 2
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(width, 12), 55)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10)
                alignment = copy(cell.alignment)
                alignment.vertical = "top"
                alignment.wrap_text = True
                cell.alignment = alignment
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-xlsx", type=Path, required=True)
    parser.add_argument("--processed-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    raw = pd.read_excel(args.raw_xlsx, sheet_name="Sheet1")
    index = pd.read_parquet(args.processed_dir / "site_index.parquet")
    audit = pd.read_csv(args.processed_dir / "site_mapping_audit.csv")
    labels = pd.read_parquet(args.processed_dir / "labels.parquet")
    tables = build_provenance_tables(raw, index, audit, labels)
    write_workbook(tables, args.output)
    print(f"Wrote {args.output}")
    print(tables["summary"].to_string())


if __name__ == "__main__":
    main()
