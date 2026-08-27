import argparse
import json
from copy import copy
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.multimodal.client_ranking import (
    LABEL_NAMES,
    WEIGHTS,
    apply_calibrators,
    fit_isotonic_calibrators,
    rank_by_balanced_confidence,
)
from src.multimodal.config import MultimodalConfig


def _ensemble_predictions(run_directories, name):
    values = [np.load(directory / f"{name}_probabilities.npy") for directory in run_directories]
    if len({value.shape for value in values}) != 1:
        raise ValueError(f"{name} prediction arrays do not share one shape")
    return np.mean(np.stack(values), axis=0)


def _annotation_text(values):
    labels = [name for name, value in zip(LABEL_NAMES, values) if int(value) == 1]
    return "; ".join(labels) if labels else "unlabeled_in_source_dataset"


def build_client_ranking(config, run_root):
    run_root = Path(run_root)
    run_directories = sorted(run_root.glob("nnpu_seed*"))
    if len(run_directories) < 2:
        raise ValueError("at least two nnPU seed directories are required")
    labels = np.load(config.processed_data / "labels.npy")
    splits = np.load(config.processed_data / "split_names.npy", allow_pickle=False)
    site_ids = np.load(config.processed_data / "dataset_site_ids.npy", allow_pickle=False)
    validation = splits == "validation"
    test = splits == "test"
    validation_raw = _ensemble_predictions(run_directories, "validation")
    test_raw = _ensemble_predictions(run_directories, "test")
    if not (len(validation_raw) == validation.sum() and len(test_raw) == test.sum()):
        raise ValueError("saved prediction rows do not match fixed split sizes")
    calibrators = fit_isotonic_calibrators(labels[validation], validation_raw)
    calibrated = apply_calibrators(calibrators, test_raw)
    ranked = rank_by_balanced_confidence(site_ids[test], calibrated)
    test_labels = labels[test]
    annotation = pd.DataFrame(
        {
            "site_id": site_ids[test],
            "source_dataset_annotation": [
                _annotation_text(values) for values in test_labels
            ],
        }
    )
    site_index = pd.read_parquet(config.processed_data / "site_index.parquet")
    ranked = ranked.merge(
        site_index[["site_id", "accession", "gene", "residue", "position"]],
        on="site_id",
        how="left",
        validate="one_to_one",
    ).merge(annotation, on="site_id", how="left", validate="one_to_one")
    ranked["customer_interpretation"] = np.where(
        ranked["source_dataset_annotation"].eq("unlabeled_in_source_dataset"),
        "High-priority model candidate; independent functional evidence is still required.",
        "The held-out source dataset contains functional annotation; this supports retrieval, not novel discovery.",
    )
    return ranked, calibrators


def _write_workbook(output_path, top10, metadata):
    workbook = Workbook()
    ranking = workbook.active
    ranking.title = "Top10_综合候选"
    note = workbook.create_sheet("说明")
    note["A1"] = "磷酸化位点多模态功能候选报告"
    note["A3"], note["B3"] = "活性调控权重", 0.40
    note["A4"], note["B4"] = "互作调控权重", 0.35
    note["A5"], note["B5"] = "稳定/降解权重", 0.25
    note["A7"] = "综合公式"
    note["B7"] = "F_balance = 0.40*p_activity + 0.35*p_interaction + 0.25*p_proteostasis"
    note["A9"] = "分数含义"
    note["B9"] = "五种子集成后，使用验证集逐任务 isotonic 校准的模型分数。"
    note["A10"] = "解释边界"
    note["B10"] = "高分表示计算优先级，不等同于湿实验已经证明功能或因果机制。"
    note["A12"] = "模型与数据"
    note["B12"] = json.dumps(metadata, ensure_ascii=False)
    headers = [
        "排名", "site_id", "活性调控置信度", "互作调控置信度", "稳定/降解置信度",
        "F_balance", "UniProt", "基因", "残基", "位置", "来源数据集注释", "客户解读",
    ]
    ranking.append(["磷酸化位点多模态功能候选 Top 10"])
    ranking.append(["按校准后分数及 F_balance=0.40A+0.35I+0.25S 排序"])
    ranking.append([])
    ranking.append(headers)
    for excel_row, (_, row) in enumerate(top10.iterrows(), start=5):
        ranking.append(
            [
                int(row["rank"]), row["site_id"], row["activity_confidence"],
                row["interaction_confidence"], row["proteostasis_confidence"], None,
                row["accession"], row["gene"], row["residue"], int(row["position"]),
                row["source_dataset_annotation"], row["customer_interpretation"],
            ]
        )
        ranking.cell(excel_row, 6).value = (
            f"=ROUND('说明'!$B$3*C{excel_row}+'说明'!$B$4*D{excel_row}+'说明'!$B$5*E{excel_row},6)"
        )
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ranking[4]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ranking[1]:
        cell.font = Font(name="Arial", bold=True, size=14, color="1F4E78")
    for sheet in (ranking, note):
        for row in sheet.iter_rows():
            for cell in row:
                font = copy(cell.font)
                font.name = "Arial"
                cell.font = font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in range(5, 15):
        for column in range(3, 7):
            ranking.cell(row, column).number_format = "0.0000"
    for column, width in {1: 8, 2: 18, 3: 16, 4: 16, 5: 18, 6: 13, 7: 14, 8: 16, 9: 8, 10: 8, 11: 28, 12: 52}.items():
        ranking.column_dimensions[get_column_letter(column)].width = width
    note.column_dimensions["A"].width = 18
    note.column_dimensions["B"].width = 100
    ranking.freeze_panes = "A5"
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Export calibrated F_balance Top 10 report")
    parser.add_argument("--config", required=True)
    parser.add_argument("--run-root", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args(argv)
    config = MultimodalConfig.load(args.config)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    ranked, calibrators = build_client_ranking(config, args.run_root)
    top10 = ranked.head(10).copy()
    metadata = {"run_root": str(args.run_root), "weights": WEIGHTS.tolist(), "calibration_fit_split": "validation", "ranking_split": "test"}
    top10.to_csv(output_dir / "client_top10_balanced_confidence.csv", index=False)
    top10.to_parquet(output_dir / "client_top10_balanced_confidence.parquet", index=False)
    ranked.to_parquet(output_dir / "client_test_ranking_balanced_confidence.parquet", index=False)
    joblib.dump(calibrators, output_dir / "isotonic_calibrators.joblib")
    (output_dir / "report_metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    _write_workbook(output_dir / "client_top10_balanced_confidence.xlsx", top10, metadata)
    print(top10.to_string(index=False))


if __name__ == "__main__":
    main()
