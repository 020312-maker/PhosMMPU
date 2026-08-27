"""Export every current-model triple-zero held-out candidate with Top30 audit fields."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.multimodal.export_triple_zero_candidates import (
    SCORE_COLUMNS,
    SEEDS,
    build_triple_zero_top_candidates,
)


def attach_top30_audit(candidates: pd.DataFrame, audit: pd.DataFrame) -> pd.DataFrame:
    """Carry manual Top30 audit values into the corresponding complete candidate rows."""
    result = candidates.copy()
    if result["位点 ID"].duplicated().any() or audit["位点 ID"].duplicated().any():
        raise ValueError("candidate and audit site IDs must be unique")
    missing = set(audit["位点 ID"].astype(str)).difference(result["位点 ID"].astype(str))
    if missing:
        raise ValueError(f"audited sites missing from complete candidates: {sorted(missing)[:5]}")
    for column in audit.columns:
        if column not in result.columns:
            result[column] = ""
    lookup = audit.set_index(audit["位点 ID"].astype(str), drop=False)
    audited_mask = result["位点 ID"].astype(str).isin(lookup.index)
    for position, site_id in result.loc[audited_mask, "位点 ID"].astype(str).items():
        row = lookup.loc[site_id]
        for column, value in row.items():
            if column not in {"候选排名", "位点 ID"} and pd.notna(value):
                result.at[position, column] = value
    result["文献核查状态"] = "未核查（仅提供检索链接）"
    result.loc[audited_mask, "文献核查状态"] = "已人工核查（原 Top30）"
    ordered = list(audit.columns) + [column for column in result.columns if column not in audit.columns]
    return result.loc[:, ordered]


def _triple_zero_count(prediction_root: Path) -> int:
    table = pd.read_parquet(Path(prediction_root) / f"nnpu_seed{SEEDS[0]}" / "test_predictions.parquet")
    return int((table[["y_activity", "y_interaction", "y_proteostasis"]].to_numpy() == 0).all(axis=1).sum())


def _format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook["全部候选"]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        heading = str(sheet.cell(1, column).value or "")
        width = min(max(len(heading) + 2, 12), 30)
        sheet.column_dimensions[letter].width = width
    note = workbook["说明"]
    note.column_dimensions["A"].width = 24
    note.column_dimensions["B"].width = 100
    for cell in note[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
    note.sheet_view.showGridLines = False
    workbook.save(path)


def export_complete_candidates(prediction_root: Path, site_index: Path, audit_xlsx: Path, output_root: Path) -> dict[str, object]:
    """Export all held-out U/U/U sites from the same prediction root as the supplied Top30."""
    output_root = Path(output_root)
    output_root.mkdir(parents=True, exist_ok=False)
    total = _triple_zero_count(prediction_root)
    candidates = build_triple_zero_top_candidates(prediction_root, site_index, total)
    audit = pd.read_excel(audit_xlsx, dtype={"位点 ID": str})
    complete = attach_top30_audit(candidates, audit)
    csv_path = output_root / "同源簇测试集_三标签全0_全部候选表.csv"
    xlsx_path = output_root / "同源簇测试集_三标签全0_全部候选表.xlsx"
    complete.to_csv(csv_path, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        complete.to_excel(writer, index=False, sheet_name="全部候选")
        pd.DataFrame([
            ("候选数量", total),
            ("候选定义", "同源簇独立测试集，三项观察标签均为 0；0 表示未标注，不表示无功能。"),
            ("排序", "0.40 x 活性调控 + 0.35 x 分子关联 + 0.25 x 稳定性/降解。"),
            ("模型", "PhosMMPU 五种子 nnPU 集成，modality dropout=0，完整网络。"),
            ("文献字段", "仅原 Top30 的文献核查字段为人工核查结果；其余候选仅提供 PubMed 检索链接，不能视为无文献证据。"),
        ], columns=["项目", "说明"]).to_excel(writer, index=False, sheet_name="说明")
    _format_workbook(xlsx_path)
    summary_path = output_root / "导出说明.json"
    summary_path.write_text(json.dumps({
        "candidate_count": total,
        "prediction_root": str(Path(prediction_root).resolve()),
        "audit_source": str(Path(audit_xlsx).resolve()),
        "audit_count": int((complete["文献核查状态"] == "已人工核查（原 Top30）").sum()),
        "scope": "全部同源簇测试集 U/U/U 候选；文献人工核查仅覆盖原 Top30。",
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"candidate_count": total, "csv": str(csv_path), "xlsx": str(xlsx_path), "summary": str(summary_path)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prediction-root", type=Path, required=True)
    parser.add_argument("--site-index", type=Path, required=True)
    parser.add_argument("--audit-xlsx", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args(argv)
    print(json.dumps(export_complete_candidates(**vars(args)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
